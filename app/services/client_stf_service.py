"""Import, validation et calculs du STF client par intervalle."""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from datetime import datetime

from datetime import date, time, timedelta
from typing import Iterable, Optional
from uuid import uuid4

from openpyxl import load_workbook
from sqlmodel import Session, select

from app.models.client_stf import ClientSTFInterval, ClientSTFPlan
from app.models.intraday import IntervalForecast
from app.services import kpi_service

INTERVAL_HOURS = 0.5


@dataclass(frozen=True)
class ClientSTFRow:
    date: date
    interval_start: time
    interval_end: time
    required_hc: float


@dataclass(frozen=True)
class ClientSTFScorecard:
    required_hc_hours: float
    scheduled_hc_hours: float
    actual_hc_hours: Optional[float]
    shortage_hc_hours: float
    surplus_hc_hours: float
    coverage_pct: float
    peak_stf_hc: float
    peak_scheduled_hc: float
    peak_actual_hc: Optional[float]
    understaffed_intervals: int
    overstaffed_intervals: int
    overtime_required_hours: float
    fte_equivalent_week: float
    model_variance_hc_hours: Optional[float]
    model_variance_pct: Optional[float]


@dataclass(frozen=True)
class EffectiveInterval:
    """Vue d'un intervalle avec le STF client comme besoin effectif."""

    interval: object
    required_hc: float

    def __getattr__(self, name: str):
        return getattr(self.interval, name)


def monday_of_week(value: date) -> date:
    return value - timedelta(days=value.weekday())


def _parse_time(value: str) -> time:
    raw = value.strip()
    if raw == "24:00":
        return time(23, 59, 59)
    parts = raw.split(":")
    if len(parts) not in (2, 3):
        raise ValueError(f"Heure invalide: {value!r}")
    hour = int(parts[0])
    minute = int(parts[1])
    second = int(parts[2]) if len(parts) == 3 else 0
    if hour > 23 or minute > 59 or second > 59:
        raise ValueError(f"Heure invalide: {value!r}")
    return time(hour, minute, second)


def parse_csv(content: str) -> list[ClientSTFRow]:
    """CSV attendu: date,interval_start,interval_end,required_hc."""
    reader = csv.DictReader(io.StringIO(content))
    required_headers = {"date", "interval_start", "interval_end", "required_hc"}
    headers = {h.strip() for h in (reader.fieldnames or [])}
    missing = required_headers - headers
    if missing:
        raise ValueError("Colonnes manquantes: " + ", ".join(sorted(missing)))

    rows: list[ClientSTFRow] = []
    seen: set[tuple[date, time]] = set()
    for line_number, raw in enumerate(reader, start=2):
        if not any((value or "").strip() for value in raw.values()):
            continue
        try:
            day = date.fromisoformat(raw["date"].strip())
            start = _parse_time(raw["interval_start"])
            end = _parse_time(raw["interval_end"])
            hc = float(raw["required_hc"])
        except (AttributeError, TypeError, ValueError) as exc:
            raise ValueError(f"Ligne {line_number} invalide: {exc}") from exc
        if hc < 0:
            raise ValueError(f"Ligne {line_number}: required_hc doit être >= 0.")
        if start == end:
            raise ValueError(f"Ligne {line_number}: intervalle vide.")
        if (day, start) in seen:
            raise ValueError(
                f"Ligne {line_number}: intervalle dupliqué pour {day} à {start:%H:%M}."
            )
        seen.add((day, start))
        rows.append(ClientSTFRow(day, start, end, hc))

    if not rows:
        raise ValueError("Le fichier STF client est vide.")
    return sorted(rows, key=lambda row: (row.date, row.interval_start))



def _row_to_stf(
    raw: dict[str, object],
    line_number: int,
    *,
    date_key: str = "date",
    start_key: str = "interval_start",
    end_key: str = "interval_end",
    hc_key: str = "required_hc",
) -> ClientSTFRow:
    try:
        raw_day = raw[date_key]
        if isinstance(raw_day, datetime):
            day = raw_day.date()
        elif isinstance(raw_day, date):
            day = raw_day
        else:
            day = date.fromisoformat(str(raw_day).strip())
        raw_start = raw[start_key]
        raw_end = raw[end_key]
        start = raw_start if isinstance(raw_start, time) else _parse_time(str(raw_start))
        end = raw_end if isinstance(raw_end, time) else _parse_time(str(raw_end))
        hc = float(raw[hc_key])
    except (AttributeError, TypeError, ValueError, KeyError) as exc:
        raise ValueError(f"Ligne {line_number} invalide: {exc}") from exc
    if hc < 0:
        raise ValueError(f"Ligne {line_number}: required_hc doit être >= 0.")
    if start == end:
        raise ValueError(f"Ligne {line_number}: intervalle vide.")
    return ClientSTFRow(day, start, end, hc)


def parse_xlsx(content: bytes) -> list[ClientSTFRow]:
    """Lit un classeur Excel: première feuille, première ligne = en-têtes."""
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration as exc:
        raise ValueError("Le fichier Excel est vide.") from exc

    header_map = {
        str(value).strip().lower(): index
        for index, value in enumerate(header_row)
        if value is not None
    }
    aliases = {
        "date": ("date", "day"),
        "interval_start": ("interval_start", "start", "heure_debut"),
        "interval_end": ("interval_end", "end", "heure_fin"),
        "required_hc": ("required_hc", "stf", "required", "hc_requis"),
    }
    indexes: dict[str, int] = {}
    for target, candidates in aliases.items():
        for candidate in candidates:
            if candidate in header_map:
                indexes[target] = header_map[candidate]
                break
        if target not in indexes:
            raise ValueError(
                f"Colonne Excel manquante pour {target}. "
                "Attendues: date, interval_start, interval_end, required_hc."
            )

    parsed: list[ClientSTFRow] = []
    seen: set[tuple[date, time]] = set()
    for line_number, values in enumerate(rows, start=2):
        if not any(value not in (None, "") for value in values):
            continue
        raw = {
            key: values[indexes[key]]
            for key in indexes
        }
        row = _row_to_stf(raw, line_number)
        key = (row.date, row.interval_start)
        if key in seen:
            raise ValueError(
                f"Ligne {line_number}: intervalle dupliqué pour "
                f"{row.date} à {row.interval_start:%H:%M}."
            )
        seen.add(key)
        parsed.append(row)

    workbook.close()
    if not parsed:
        raise ValueError("Le fichier Excel STF client est vide.")
    return sorted(parsed, key=lambda row: (row.date, row.interval_start))


def create_plan(
    session: Session,
    *,
    week_start_date: date,
    campaign_id: int,
    skill_id: int,
    rows: Iterable[ClientSTFRow],
    label: str = "STF client",
    notes: str | None = None,
    created_by_user_id: int | None = None,
) -> ClientSTFPlan:
    week_start_date = monday_of_week(week_start_date)
    materialized = list(rows)
    if not materialized:
        raise ValueError("Aucune ligne STF à importer.")

    week_end = week_start_date + timedelta(days=6)
    seen: set[tuple[date, time]] = set()
    for row in materialized:
        if not (week_start_date <= row.date <= week_end):
            raise ValueError(
                f"{row.date} n'appartient pas à la semaine {week_start_date}."
            )
        key = (row.date, row.interval_start)
        if key in seen:
            raise ValueError(
                f"Intervalle dupliqué: {row.date} {row.interval_start:%H:%M}."
            )
        seen.add(key)
        if row.required_hc < 0:
            raise ValueError("required_hc doit être >= 0.")

    batch_id = uuid4().hex
    current = session.exec(
        select(ClientSTFPlan).where(
            ClientSTFPlan.week_start_date == week_start_date,
            ClientSTFPlan.campaign_id == campaign_id,
            ClientSTFPlan.skill_id == skill_id,
            ClientSTFPlan.is_current == True,  # noqa: E712
        )
    ).all()
    for plan in current:
        plan.is_current = False
        session.add(plan)

    plan = ClientSTFPlan(
        week_start_date=week_start_date,
        campaign_id=campaign_id,
        skill_id=skill_id,
        label=label or "STF client",
        notes=notes,
        created_by=created_by_user_id,
        import_batch_id=batch_id,
        is_current=True,
    )
    session.add(plan)
    session.flush()

    for row in materialized:
        session.add(
            ClientSTFInterval(
                plan_id=plan.id,
                date=row.date,
                interval_start=row.interval_start,
                interval_end=row.interval_end,
                required_hc=row.required_hc,
            )
        )

    session.commit()
    session.refresh(plan)
    return plan


def current_plan(
    session: Session,
    *,
    target_date: date,
    campaign_id: int,
    skill_id: int,
) -> ClientSTFPlan | None:
    return session.exec(
        select(ClientSTFPlan).where(
            ClientSTFPlan.week_start_date == monday_of_week(target_date),
            ClientSTFPlan.campaign_id == campaign_id,
            ClientSTFPlan.skill_id == skill_id,
            ClientSTFPlan.is_current == True,  # noqa: E712
        )
    ).first()


def list_intervals(session: Session, plan_id: int) -> list[ClientSTFInterval]:
    return list(
        session.exec(
            select(ClientSTFInterval)
            .where(ClientSTFInterval.plan_id == plan_id)
            .order_by(ClientSTFInterval.date, ClientSTFInterval.interval_start)
        ).all()
    )


def effective_intervals(
    intervals: list[IntervalForecast],
    client_rows: list[ClientSTFInterval],
) -> list[EffectiveInterval]:
    by_key = {(row.date, row.interval_start): row.required_hc for row in client_rows}
    return [
        EffectiveInterval(
            interval=interval,
            required_hc=by_key.get((interval.date, interval.interval_start), interval.required_hc),
        )
        for interval in intervals
    ]



def effective_intervals_for_range(
    session: Session,
    intervals: list[IntervalForecast],
    *,
    campaign_id: int,
    skill_id: int,
) -> list[EffectiveInterval]:
    """Applique le STF client courant à tous les intervalles d'une plage."""
    if not intervals:
        return []
    week_starts = sorted({monday_of_week(interval.date) for interval in intervals})
    rows: list[ClientSTFInterval] = []
    for week_start in week_starts:
        plan = session.exec(
            select(ClientSTFPlan).where(
                ClientSTFPlan.week_start_date == week_start,
                ClientSTFPlan.campaign_id == campaign_id,
                ClientSTFPlan.skill_id == skill_id,
                ClientSTFPlan.is_current == True,  # noqa: E712
            )
        ).first()
        if plan is not None:
            rows.extend(list_intervals(session, plan.id))
    return effective_intervals(intervals, rows)


def scorecard(
    intervals: list[IntervalForecast | EffectiveInterval],
    *,
    client_rows: list[ClientSTFInterval],
    weekly_contract_hours: float = 40.0,
) -> ClientSTFScorecard:
    row_by_key = {(row.date, row.interval_start): row.required_hc for row in client_rows}
    matched = [
        (interval, row_by_key[(interval.date, interval.interval_start)])
        for interval in intervals
        if (interval.date, interval.interval_start) in row_by_key
    ]
    required = sum(max(stf_hc, 0) * INTERVAL_HOURS for _, stf_hc in matched)
    scheduled = sum(max(interval.scheduled_hc, 0) * INTERVAL_HOURS for interval, _ in matched)
    actual_values = [interval.actual_hc for interval, _ in matched if interval.actual_hc is not None]
    actual = sum(actual_values) * INTERVAL_HOURS if actual_values else None
    shortage = sum(
        max(stf_hc - max(interval.scheduled_hc, 0), 0) * INTERVAL_HOURS
        for interval, stf_hc in matched
    )
    surplus = sum(
        max(max(interval.scheduled_hc, 0) - stf_hc, 0) * INTERVAL_HOURS
        for interval, stf_hc in matched
    )
    covered = sum(
        min(max(stf_hc, 0), max(interval.scheduled_hc, 0)) * INTERVAL_HOURS
        for interval, stf_hc in matched
    )
    peak_stf = max((stf_hc for _, stf_hc in matched), default=0)
    peak_scheduled = max((interval.scheduled_hc for interval, _ in matched), default=0)
    peak_actual = max(actual_values) if actual_values else None

    model_variance_hours: float | None = None
    model_variance_pct: float | None = None
    model_pairs = [
        (
            stf_hc,
            interval.interval.required_hc
            if isinstance(interval, EffectiveInterval)
            else interval.required_hc,
        )
        for interval, stf_hc in matched
    ]
    if model_pairs:
        model_variance_hours = sum((stf - model) * INTERVAL_HOURS for stf, model in model_pairs)
        model_total = sum(max(model, 0) * INTERVAL_HOURS for _, model in model_pairs)
        if model_total:
            model_variance_pct = model_variance_hours / model_total * 100

    return ClientSTFScorecard(
        required_hc_hours=required,
        scheduled_hc_hours=scheduled,
        actual_hc_hours=actual,
        shortage_hc_hours=shortage,
        surplus_hc_hours=surplus,
        coverage_pct=kpi_service.coverage_pct(required, covered),
        peak_stf_hc=peak_stf,
        peak_scheduled_hc=peak_scheduled,
        peak_actual_hc=peak_actual,
        understaffed_intervals=sum(
            1 for interval, stf_hc in matched if stf_hc - interval.scheduled_hc > 0.5
        ),
        overstaffed_intervals=sum(
            1 for interval, stf_hc in matched if interval.scheduled_hc - stf_hc > 0.5
        ),
        overtime_required_hours=shortage,
        fte_equivalent_week=required / weekly_contract_hours if weekly_contract_hours > 0 else 0,
        model_variance_hc_hours=model_variance_hours,
        model_variance_pct=model_variance_pct,
    )
